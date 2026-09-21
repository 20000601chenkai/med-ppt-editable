import copy
import json
import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from build_deck import build, validate
from audit_pptx import audit


class DeckTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.output = self.folder / "test.pptx"
        self.data = json.loads((ROOT / "examples" / "case-demo.json").read_text(encoding="utf-8"))

    def test_native_objects_data_and_gap_survive_save(self):
        manifest = build(self.data, self.output)
        report = audit(self.output, manifest)
        self.assertEqual(report["errors"], [])
        self.assertEqual(sum(s["tables"] for s in report["slides"]), 1)
        self.assertEqual(sum(s["workbooks"] for s in report["slides"]), 1)
        prs = Presentation(self.output)
        chart = next(s.chart for s in prs.slides[4].shapes if s.has_chart)
        self.assertEqual(list(chart.series[0].values), [20, 16, None, 9])
        self.assertEqual(chart._chartSpace.xpath(".//c:dispBlanksAs")[0].get("val"), "gap")

    def test_text_table_and_chart_can_be_changed_and_reopened(self):
        manifest = build(self.data, self.output)
        prs = Presentation(self.output)
        title = next(s for s in prs.slides[0].shapes if s.name == "cover--title")
        title.text = "Revised title"
        table = next(s.table for s in prs.slides[3].shapes if s.has_table)
        table.cell(1, 1).text = "21"
        chart = next(s.chart for s in prs.slides[4].shapes if s.has_chart)
        data = CategoryChartData()
        data.categories = self.data["slides"][4]["categories"]
        data.add_series("Revised series", [21, 17, None, 8])
        chart.replace_data(data)
        revised = self.folder / "revised.pptx"
        prs.save(revised)
        reopened = Presentation(revised)
        self.assertEqual(next(s for s in reopened.slides[0].shapes if s.name == "cover--title").text, "Revised title")
        self.assertEqual(next(s.table for s in reopened.slides[3].shapes if s.has_table).cell(1, 1).text, "21")
        self.assertEqual(audit(revised)["errors"], [])
        self.assertFalse(audit(revised, manifest)["structural_ok"])

    def test_invalid_source_numeric_and_shape_data_fail(self):
        for mutator in (
            lambda d: d["slides"][0].update(source_ids=["MISSING"]),
            lambda d: d["slides"][4]["series"][0].update(values=[1, 2]),
            lambda d: d["slides"][4]["series"][0].update(values=[1, True, 2, 3]),
            lambda d: d["slides"][4]["series"][0].update(values=[1, float("nan"), 2, 3]),
            lambda d: d["slides"][0].update(unknown_field="typo"),
            lambda d: d["slides"][1].update(id="cover"),
            lambda d: d["slides"][3]["rows"].append(["short"]),
        ):
            changed = copy.deepcopy(self.data)
            mutator(changed)
            with self.assertRaises(ValueError):
                validate(changed)

    def test_overflow_and_overwrite_are_rejected(self):
        self.data["slides"][0]["title"] = "Long title " * 11
        with self.assertRaises(ValueError):
            build(self.data, self.output)
        self.assertFalse(self.output.exists())
        self.data["slides"][0]["title"] = "Short title"
        build(self.data, self.output)
        before = self.output.read_bytes()
        with self.assertRaises(ValueError):
            build(self.data, self.output)
        self.assertEqual(self.output.read_bytes(), before)

    def test_image_preserves_aspect_ratio_and_alt_text(self):
        image = self.folder / "image.png"
        Image.new("RGB", (900, 300), "#087E8B").save(image)
        self.data["slides"] = [{"id": "image", "layout": "image", "title": "Figure", "image": "image.png", "caption": "Synthetic test image", "alt": "A solid color rectangle"}]
        manifest = build(self.data, self.output, self.folder)
        self.assertEqual(audit(self.output, manifest)["errors"], [])
        shape = next(s for s in Presentation(self.output).slides[0].shapes if s.name == "image--image")
        self.assertAlmostEqual(shape.width / shape.height, 3, places=4)
        self.assertEqual(shape.image.blob, image.read_bytes())
        self.assertEqual(shape._element.nvPicPr.cNvPr.get("descr"), "A solid color rectangle")

    def test_negative_bars_and_missing_values(self):
        self.data["slides"] = [self.data["slides"][4]]
        self.data["slides"][0]["chart_type"] = "bar"
        self.data["slides"][0]["series"][0]["values"] = [-3, 0, None, 8]
        manifest = build(self.data, self.output)
        self.assertEqual(audit(self.output, manifest)["errors"], [])
        chart = next(s.chart for s in Presentation(self.output).slides[0].shapes if s.has_chart)
        self.assertEqual(chart.value_axis.minimum_scale, -3)

    def test_out_of_bounds_fails_and_large_picture_warns(self):
        build(self.data, self.output)
        prs = Presentation(self.output)
        image = self.folder / "background.png"
        Image.new("RGB", (160, 90), "white").save(image)
        prs.slides[0].shapes.add_picture(str(image), 0, 0, prs.slide_width, prs.slide_height)
        prs.slides[1].shapes[1].left = Inches(-1)
        revised = self.folder / "bounds.pptx"
        prs.save(revised)
        result = audit(revised)
        self.assertTrue(any("out-of-bounds" in e for e in result["errors"]))
        self.assertTrue(any("picture covers" in w for w in result["warnings"]))

    def test_journal_club_scaffold_builds(self):
        data = json.loads((ROOT / "examples" / "journal-club.json").read_text(encoding="utf-8"))
        manifest = build(data, self.output)
        self.assertEqual(audit(self.output, manifest)["errors"], [])


if __name__ == "__main__":
    unittest.main()
