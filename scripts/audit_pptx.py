"""Inspect native PPTX objects; structural checks do not replace Office rendering."""

import argparse
import io
import json
import sys
from pathlib import Path
from zipfile import ZipFile, BadZipFile
import xml.etree.ElementTree as ET

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def shapes_in(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from shapes_in(shape.shapes)


def workbook_rows(blob):
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(io.BytesIO(blob)) as book:
        shared = []
        if "xl/sharedStrings.xml" in book.namelist():
            shared = ["".join(si.itertext()) for si in ET.fromstring(book.read("xl/sharedStrings.xml"))]
        root = ET.fromstring(book.read("xl/worksheets/sheet1.xml"))
        cells = {}
        for cell in root.findall(".//s:sheetData/s:row/s:c", ns):
            value = cell.find("s:v", ns)
            if cell.get("t") == "inlineStr":
                parsed = "".join(cell.find("s:is", ns).itertext())
            elif value is None:
                parsed = None
            elif cell.get("t") == "s":
                parsed = shared[int(value.text)]
            else:
                parsed = float(value.text)
            cells[cell.get("r")] = parsed
        return cells


def audit(path, manifest=None):
    with ZipFile(path) as package:
        bad = package.testzip()
        if bad:
            raise ValueError(f"Corrupt ZIP member: {bad}")
    prs = Presentation(str(path))
    report = {"structural_ok": True, "visual_check": "not_performed_by_this_script", "errors": [], "warnings": [], "slides": []}
    errors, warnings = report["errors"], report["warnings"]
    if manifest is not None:
        if manifest.get("schema_version") != 1 or not isinstance(manifest.get("slides"), list):
            raise ValueError("Unsupported or malformed manifest")
        if len(manifest["slides"]) != len(prs.slides):
            errors.append("Slide count differs from manifest")
    for si, slide in enumerate(prs.slides):
        row = {"number": si + 1, "name": slide.name, "text_frames": 0, "tables": 0, "charts": 0, "pictures": 0, "workbooks": 0}
        texts, tables, charts = [], [], []
        for shape in slide.shapes:
            if shape.left < -9144 or shape.top < -9144 or shape.left + shape.width > prs.slide_width + 9144 or shape.top + shape.height > prs.slide_height + 9144:
                errors.append(f"Slide {si+1}: out-of-bounds object {shape.name}")
        for shape in shapes_in(slide.shapes):
            if shape.has_text_frame and shape.text.strip():
                row["text_frames"] += 1
                texts.extend(shape.text.splitlines())
            if shape.has_table:
                row["tables"] += 1
                tables.append([[cell.text for cell in tr.cells] for tr in shape.table.rows])
            if shape.has_chart:
                row["charts"] += 1
                chart = shape.chart
                actual = {"categories": [c.label for c in chart.plots[0].categories],
                          "series": [{"name": s.name, "values": list(s.values)} for s in chart.series], "type": int(chart.chart_type)}
                charts.append(actual)
                books = [rel.target_part for rel in chart.part.rels.values() if rel.reltype.endswith("/package") and not rel.is_external]
                if len(books) != 1:
                    errors.append(f"Slide {si+1}: chart has no single embedded workbook")
                else:
                    try:
                        cells = workbook_rows(books[0].blob)
                        row["workbooks"] += 1
                        for ci, category in enumerate(actual["categories"], 2):
                            if cells.get(f"A{ci}") != category:
                                errors.append(f"Slide {si+1}: chart workbook category differs from chart cache")
                        for ni, series in enumerate(actual["series"], 1):
                            column = chr(ord("A") + ni)
                            if cells.get(f"{column}1") != series["name"]:
                                errors.append(f"Slide {si+1}: chart workbook series name differs from chart cache")
                            for vi, value in enumerate(series["values"], 2):
                                if cells.get(f"{column}{vi}") != value:
                                    errors.append(f"Slide {si+1}: chart workbook data differs from chart cache")
                    except (ValueError, KeyError, BadZipFile, ET.ParseError, IndexError) as exc:
                        errors.append(f"Slide {si+1}: cannot verify chart workbook: {exc}")
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                row["pictures"] += 1
                area = shape.width * shape.height / (prs.slide_width * prs.slide_height)
                if area > .70:
                    warnings.append(f"Slide {si+1}: picture covers {area:.0%} of slide; review flattening and intended content")
        if not texts and not tables and not charts:
            warnings.append(f"Slide {si+1}: no native text, table, or chart found")
        if manifest is not None and si < len(manifest["slides"]):
            expected = manifest["slides"][si]
            if slide.name != expected["id"]:
                errors.append(f"Slide {si+1}: ID/order differs from manifest")
            for value in expected["text"]:
                if value and value not in texts:
                    errors.append(f"Slide {si+1}: missing expected editable text: {value[:70]}")
            if tables != expected["tables"]:
                errors.append(f"Slide {si+1}: native table count or cell data differs from manifest")
            if charts != expected["charts"]:
                errors.append(f"Slide {si+1}: native chart count, type, or data differs from manifest")
            if row["pictures"] != expected["images"]:
                errors.append(f"Slide {si+1}: picture count differs from manifest")
        report["slides"].append(row)
    report["structural_ok"] = not errors
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--manifest", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    try:
        if args.output and args.output.resolve() in {p.resolve() for p in (args.input, args.manifest) if p}:
            raise ValueError("Report output must not overwrite input or manifest")
        manifest = json.loads(args.manifest.read_text(encoding="utf-8-sig")) if args.manifest else None
        result = audit(args.input, manifest)
        payload = json.dumps(result, ensure_ascii=False, indent=2)
        if args.output:
            args.output.parent.mkdir(parents=True, exist_ok=True)
            args.output.write_text(payload, encoding="utf-8")
        else:
            print(payload)
        print(f"Slides: {len(result['slides'])}; errors: {len(result['errors'])}; warnings: {len(result['warnings'])}")
        return 0 if result["structural_ok"] else 1
    except (ValueError, OSError, KeyError, TypeError, BadZipFile, ET.ParseError) as exc:
        print(f"Audit failed: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
