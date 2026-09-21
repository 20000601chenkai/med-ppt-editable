"""Build a small medical slide deck using native PowerPoint objects."""

import argparse
import json
import math
import re
import sys
import unicodedata
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


W, H = 13.333333, 7.5
INK, MUTED, GREEN, RED, PALE = "20282A", "626D72", "087E8B", "BA4058", "F0F4F4"
COMMON = {"id", "layout", "title", "subtitle", "notes", "source_ids"}
LAYOUTS = {
    "cover": ({"body"}, {"body"}),
    "bullets": ({"items"}, {"items"}),
    "timeline": ({"events"}, {"events"}),
    "table": ({"columns", "rows"}, {"columns", "rows"}),
    "chart": ({"categories", "series", "chart_type", "y_title", "takeaway"}, {"categories", "series"}),
    "image": ({"image", "caption", "alt"}, {"image", "caption", "alt"}),
}


def keys(value, allowed, required, label):
    if not isinstance(value, dict):
        raise ValueError(f"{label}: expected an object")
    unknown, missing = set(value) - allowed, required - set(value)
    if unknown or missing:
        raise ValueError(f"{label}: unknown fields {sorted(unknown)}; missing {sorted(missing)}")


def string(value, label, maximum=4000):
    if not isinstance(value, str) or not value.strip() or len(value) > maximum:
        raise ValueError(f"{label}: expected a nonempty string of at most {maximum} characters")
    if any(ord(c) < 32 and c not in "\n\t\r" for c in value):
        raise ValueError(f"{label}: XML control characters are not allowed")


def sequence(value, label, low, high):
    if not isinstance(value, list) or not low <= len(value) <= high:
        raise ValueError(f"{label}: expected {low}-{high} items")


def validate(data):
    keys(data, {"title", "subtitle", "demo", "font", "sources", "slides"},
         {"title", "sources", "slides"}, "deck")
    string(data["title"], "title", 120)
    for field in ("subtitle", "font"):
        if field in data:
            string(data[field], field, 160)
    if "demo" in data and not isinstance(data["demo"], bool):
        raise ValueError("demo: expected a boolean")
    if not isinstance(data["sources"], dict):
        raise ValueError("sources: expected an object")
    for sid, citation in data["sources"].items():
        string(sid, "source ID", 30)
        string(citation, f"source {sid}", 8000)
    sequence(data["slides"], "slides", 1, 100)
    ids = set()
    for s in data["slides"]:
        if not isinstance(s, dict) or not isinstance(s.get("layout"), str) or s["layout"] not in LAYOUTS:
            raise ValueError("slide: missing or unsupported layout")
        allowed, required = LAYOUTS[s["layout"]]
        keys(s, COMMON | allowed, {"id", "title", "layout"} | required, "slide")
        string(s["id"], "slide ID", 80)
        if not re.fullmatch(r"[A-Za-z0-9_-]+", s["id"]) or s["id"] in ids:
            raise ValueError(f"Invalid or duplicate slide ID: {s['id']}")
        ids.add(s["id"])
        for field in ("title", "subtitle", "notes", "body", "y_title", "takeaway", "image", "caption", "alt"):
            if field in s:
                string(s[field], f"{s['id']}.{field}")
        refs = s.get("source_ids", [])
        sequence(refs, "source_ids", 0, 20)
        for ref in refs:
            string(ref, "source ID", 30)
            if ref not in data["sources"]:
                raise ValueError(f"Unknown source ID: {ref}")
        if s["layout"] == "bullets":
            sequence(s["items"], "items", 1, 6)
            for item in s["items"]:
                string(item, "bullet")
        if s["layout"] == "timeline":
            sequence(s["events"], "events", 2, 5)
            for event in s["events"]:
                keys(event, {"label", "text"}, {"label", "text"}, "event")
                string(event["label"], "event.label")
                string(event["text"], "event.text")
        if s["layout"] == "table":
            sequence(s["columns"], "columns", 2, 5)
            sequence(s["rows"], "rows", 1, 8)
            for row in [s["columns"]] + s["rows"]:
                sequence(row, "table row", len(s["columns"]), len(s["columns"]))
                for cell in row:
                    string(cell, "table cell")
        if s["layout"] == "chart":
            if s.get("chart_type", "line") not in ("line", "bar"):
                raise ValueError("chart_type: expected line or bar")
            sequence(s["categories"], "categories", 2, 12)
            for category in s["categories"]:
                string(category, "category", 24)
            if len(set(s["categories"])) != len(s["categories"]):
                raise ValueError("chart categories must be unique")
            sequence(s["series"], "series", 1, 4)
            for series in s["series"]:
                keys(series, {"name", "values"}, {"name", "values"}, "series")
                string(series["name"], "series.name", 60)
                sequence(series["values"], "series.values", len(s["categories"]), len(s["categories"]))
                if all(v is None for v in series["values"]):
                    raise ValueError("a chart series needs at least one observed value")
                for value in series["values"]:
                    if value is not None and (type(value) not in (int, float) or not math.isfinite(value)):
                        raise ValueError("chart values must be finite numbers or null")
    return data


def check_space(text, width, height, size, label):
    """Conservative script-aware estimate; actual Office rendering still needs review."""
    capacity = max(1, (width * 72 - 8) / size)
    lines = 0
    for paragraph in text.split("\n"):
        units = sum(1.0 if unicodedata.east_asian_width(c) in "WF" else 0.60 for c in paragraph)
        lines += max(1, math.ceil(units / capacity))
    if lines * size * 1.3 > height * 72 - 4:
        raise ValueError(f"Text may overflow at {label}; shorten or split the slide")


def style_run(run, font, size, color, bold=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    props = run._r.get_or_add_rPr()
    east = props.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if east is None:
        east = OxmlElement("a:ea")
        props.append(east)
    east.set("typeface", font)


def text_frame(frame, text, font, size, color=INK, bold=False):
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(text.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        p.line_spacing = 1.15
        run = p.add_run()
        run.text = line
        style_run(run, font, size, color, bold)


def build(data, output, base_dir=None):
    validate(data)
    output = Path(output).resolve()
    base_dir = Path(base_dir or ".").resolve()
    if output.suffix.lower() != ".pptx":
        raise ValueError("output must end with .pptx")
    if output.exists() or output.with_suffix(".manifest.json").exists():
        raise ValueError("Output or manifest already exists; choose a new filename")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.title = data["title"]
    prs.core_properties.author = ""
    prs.core_properties.last_modified_by = ""
    prs.core_properties.subject = "Synthetic demonstration" if data.get("demo") else "Medical presentation"
    font = data.get("font", "Microsoft YaHei")
    manifest = {"schema_version": 1, "slides": []}
    for index, spec in enumerate(data["slides"], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.name = spec["id"]
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        expected = {"id": spec["id"], "layout": spec["layout"], "text": [], "tables": [], "charts": [], "images": 0}

        def box(key, value, x, y, w, h, size=22, color=INK, bold=False):
            check_space(value, w, h, size, f"{spec['id']}.{key}")
            shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
            shape.name = f"{spec['id']}--{key}"
            shape.text_frame.margin_left = shape.text_frame.margin_right = 0
            shape.text_frame.margin_top = shape.text_frame.margin_bottom = 0
            text_frame(shape.text_frame, value, font, size, color, bold)
            expected["text"].extend(value.split("\n"))
            return shape

        def block(key, x, y, w, h, color=GREEN, kind=MSO_SHAPE.RECTANGLE):
            shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
            shape.name = f"{spec['id']}--{key}"
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(color)
            shape.line.fill.background()
            shape._element.spPr.append(OxmlElement("a:effectLst"))
            return shape

        block("accent", .65, .42, .65, .06)
        if spec["layout"] == "cover":
            box("title", spec["title"], .8, 2.0, 11.7, 1.6, 40, bold=True)
            if spec.get("subtitle"):
                box("subtitle", spec["subtitle"], .82, 3.8, 11.5, .6, 22, GREEN)
            box("body", spec["body"], .82, 5.0, 11.5, 1.1, 18, MUTED)
            block("cover-mark", 11.9, .88, .4, .4, RED, MSO_SHAPE.OVAL)
        else:
            box("title", spec["title"], .65, .68, 12.0, .75, 30, bold=True)
            if spec.get("subtitle"):
                box("subtitle", spec["subtitle"], .68, 1.5, 11.95, .42, 15, MUTED)
            if spec["layout"] == "bullets":
                step = 4.65 / len(spec["items"])
                for i, item in enumerate(spec["items"]):
                    box(f"number-{i}", f"{i+1:02d}", .7, 2.15 + i * step, .52, .55, 17, GREEN, True)
                    box(f"item-{i}", item, 1.4, 2.12 + i * step, 11.0, step - .1, 22)
            elif spec["layout"] == "timeline":
                events = spec["events"]
                span = 12.0 / len(events)
                line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(.85), Inches(3.35), Inches(12.3), Inches(3.35))
                line.name = f"{spec['id']}--connector"
                line.line.color.rgb = RGBColor.from_string(GREEN)
                line.line.width = Pt(2)
                for i, event in enumerate(events):
                    x = .75 + i * span
                    box(f"event-label-{i}", event["label"], x, 2.4, span - .22, .65, 20, GREEN, True)
                    block(f"event-node-{i}", x, 3.25, .2, .2, GREEN, MSO_SHAPE.OVAL)
                    box(f"event-text-{i}", event["text"], x, 3.9, span - .25, 2.2, 20)
            elif spec["layout"] == "table":
                rows = [spec["columns"]] + spec["rows"]
                height = max(3.25, len(rows) * .52)
                frame = slide.shapes.add_table(len(rows), len(rows[0]), Inches(.65), Inches(2.15), Inches(12.0), Inches(height))
                frame.name = f"{spec['id']}--table"
                for ri, row in enumerate(rows):
                    for ci, value in enumerate(row):
                        cell = frame.table.cell(ri, ci)
                        cell.margin_left = cell.margin_right = Inches(.1)
                        cell.margin_top = cell.margin_bottom = Inches(.07)
                        check_space(value, 12.0 / len(row) - .2, height / len(rows) - .14, 17, f"table[{ri},{ci}]")
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor.from_string(GREEN if ri == 0 else (PALE if ri % 2 else "FFFFFF"))
                        text_frame(cell.text_frame, value, font, 17, "FFFFFF" if ri == 0 else INK, ri == 0)
                expected["tables"].append(rows)
            elif spec["layout"] == "chart":
                chart_data = CategoryChartData()
                chart_data.categories = spec["categories"]
                for series in spec["series"]:
                    chart_data.add_series(series["name"], series["values"])
                kind = XL_CHART_TYPE.LINE_MARKERS if spec.get("chart_type", "line") == "line" else XL_CHART_TYPE.COLUMN_CLUSTERED
                frame = slide.shapes.add_chart(kind, Inches(.85), Inches(2.12), Inches(11.65), Inches(4.12), chart_data)
                frame.name = f"{spec['id']}--chart"
                chart = frame.chart
                chart.font.name = font
                chart.font.size = Pt(14)
                chart.has_title = False
                # python-pptx 1.0 does not expose this setting as a public property.
                for blanks in chart._chartSpace.xpath(".//c:dispBlanksAs"):
                    blanks.set("val", "gap")
                chart.has_legend = len(spec["series"]) > 1
                if chart.has_legend:
                    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                    chart.legend.include_in_layout = False
                chart.value_axis.has_major_gridlines = True
                chart.value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string("DDE4E6")
                chart.value_axis.major_gridlines.format.line.width = Pt(.6)
                for axis in (chart.category_axis, chart.value_axis):
                    axis.tick_labels.font.name = font
                    axis.tick_labels.font.size = Pt(14)
                    axis.format.line.color.rgb = RGBColor.from_string("A8B4B8")
                if kind == XL_CHART_TYPE.COLUMN_CLUSTERED:
                    chart.value_axis.minimum_scale = min(0, min(v for item in spec["series"] for v in item["values"] if v is not None))
                if spec.get("y_title"):
                    chart.value_axis.has_title = True
                    text_frame(chart.value_axis.axis_title.text_frame, spec["y_title"], font, 14, MUTED)
                for si, series in enumerate(chart.series):
                    color = [GREEN, RED, "466BD7", "717C48"][si]
                    series.format.line.color.rgb = RGBColor.from_string(color)
                    series.format.line.width = Pt(2)
                    if kind == XL_CHART_TYPE.COLUMN_CLUSTERED:
                        series.format.fill.solid()
                        series.format.fill.fore_color.rgb = RGBColor.from_string(color)
                    else:
                        series.marker.style = XL_MARKER_STYLE.CIRCLE
                        series.marker.size = 6
                        series.marker.format.fill.solid()
                        series.marker.format.fill.fore_color.rgb = RGBColor.from_string(color)
                        series.marker.format.line.color.rgb = RGBColor.from_string(color)
                for props in chart._chartSpace.xpath(".//a:defRPr"):
                    east = props.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
                    if east is None:
                        east = OxmlElement("a:ea")
                        props.append(east)
                    east.set("typeface", font)
                if spec.get("takeaway"):
                    box("takeaway", spec["takeaway"], .85, 6.35, 11.65, .48, 15, MUTED)
                expected["charts"].append({"categories": spec["categories"], "series": spec["series"], "type": int(kind)})
            elif spec["layout"] == "image":
                path = (base_dir / spec["image"]).resolve()
                with Image.open(path) as picture:
                    if picture.format not in ("PNG", "JPEG"):
                        raise ValueError("image must be PNG or JPEG")
                    ratio = min(11.9 / picture.width, 3.85 / picture.height)
                    iw, ih = picture.width * ratio, picture.height * ratio
                shape = slide.shapes.add_picture(str(path), Inches((W-iw)/2), Inches(2.1+(3.85-ih)/2), Inches(iw), Inches(ih))
                shape.name = f"{spec['id']}--image"
                shape._element.nvPicPr.cNvPr.set("descr", spec["alt"])
                box("caption", spec["caption"], .75, 6.12, 11.85, .65, 16, MUTED)
                expected["images"] = 1

        label = "\u6f14\u793a\u6570\u636e / \u975e\u4e34\u5e8a\u8bc1\u636e" if data.get("demo") else data["title"]
        box("footer", label, .68, 7.06, 7.0, .26, 10, RED if data.get("demo") else MUTED)
        refs = spec.get("source_ids", [])
        if refs:
            box("sources", " / ".join(refs), 8.0, 7.06, 3.5, .26, 10, MUTED)
        box("page", f"{index:02d} / {len(data['slides']):02d}", 11.9, 7.04, .85, .28, 10, MUTED)
        notes = [spec.get("notes", "")]
        notes.extend(f"[{ref}] {data['sources'][ref]}" for ref in refs)
        slide.notes_slide.notes_text_frame.text = "\n\n".join(n for n in notes if n)
        manifest["slides"].append(expected)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    output.with_suffix(".manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    try:
        data = json.loads(args.input.read_text(encoding="utf-8-sig"))
        result = build(data, args.output, args.input.parent)
    except (ValueError, OSError, KeyError) as exc:
        print(f"Build failed: {exc}", file=sys.stderr)
        return 2
    print(f"Built {len(result['slides'])} slides: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
